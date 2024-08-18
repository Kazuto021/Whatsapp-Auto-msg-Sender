import pyautogui as pag
import openpyxl as oxl
import time as t 
# import re 

def num_extractor():
    li=[]
    phone_index = 0
    path = "demoh.xlsx"
    curr_wb = oxl.load_workbook(path)
    file = curr_wb.active
    row = file.max_row
    col = file.max_column
    print("Total rows:",row)
    print("Total columns:",col)
    for i in range(1,col+1):
        cell_object = file.cell(row=1,column = i)
        print(cell_object.value)
        if cell_object.value.lower() == "phone":
            phone_index = i
            print("Phone column found at:",phone_index)
            break   
    for j in range(2,row+1):
        cell_object = file.cell(row=j,column = phone_index)
        li.append(cell_object.value)
    return li

def auto_send(text):
    pag.typewrite(text,interval=.100)
    # pag.press("enter")
    pag.write("\n")  

def search_msg(text):
    li=num_extractor()
    t.sleep(5)
    ready=pag.prompt("are you ready?")
    if ready.lower()=="yes" or "haan":
        for i in li:
            # print(i)
            pag.press("tab",7)
            t.sleep(1)
            pag.typewrite(str(i),interval=.2)
            pag.sleep(2)
            pag.press("Enter")
            pag.click(1186,1012)
            pag.sleep(1)
            auto_send(text)
            pag.press("tab")

__inp__ = str(pag.prompt("Enter the Text:\n"))
# li= str(__inp__).split(",")
# text=li[0]
search_msg(__inp__)
# pag.sleep(5)
# pag.typewrite("Trial message.\n")
# pag.press("enter")
